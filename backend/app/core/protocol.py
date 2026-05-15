import json
import logging
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .generation import format_context
from .llm import BaseLLMProvider

logger = logging.getLogger(__name__)

PROTOCOL_SYSTEM_PROMPT = """\
You are an expert clinical research physician who designs clinical trial and
real-world evidence study protocols. Given a set of reference trials retrieved
from ClinicalTrials.gov, draft a NEW breast cancer study protocol.

First, INFER THE STUDY TYPE from the user's query:
- If the query mentions "real-world evidence", "RWE", "retrospective", "cohort",
  "registry", "claims data", "EHR-based", or "observational", treat this as an
  OBSERVATIONAL / REAL-WORLD EVIDENCE study — NOT a phase-numbered interventional
  trial. Set "phase": "N/A" (observational studies do not have phases).
- If the query specifies "Phase I", "Phase II", "Phase III", or "Phase IV", use
  that phase.
- If no study type is specified, default to "Phase II" interventional.

Return your answer as a single JSON object (no markdown fences) with exactly
these keys:

{
  "title": "Short descriptive trial title",
  "official_title": "Full formal title including phase and design",
  "protocol_id": "",
  "sponsor": "",
  "phase": "Phase II (or Phase I / III / IV; use 'N/A' for observational / RWE)",
  "status": "PLANNED",
  "conditions": "Specific breast cancer subtype(s)",
  "summary": "2-3 sentence plain-language summary",
  "description": "Detailed scientific background and rationale (2-3 paragraphs covering unmet need, mechanism of action, preclinical/Phase I evidence)",
  "primary_objectives": ["objective 1", "..."],
  "secondary_objectives": ["objective 1", "..."],
  "exploratory_objectives": ["objective 1", "..."],
  "study_design": "Randomized, Open-label, ...",
  "study_schema": "Description of study arms, randomization ratio, stratification factors, and overall study flow",
  "intervention_name": "Drug(s) / therapy name(s)",
  "intervention_description": "Detailed dosing, schedule, route of administration, dose modifications, and treatment arms",
  "comparator": "Comparator arm description (or 'Single-arm, no comparator')",
  "treatment_duration": "Duration of treatment and follow-up periods",
  "inclusion_criteria": ["criterion 1", "criterion 2", "..."],
  "exclusion_criteria": ["criterion 1", "criterion 2", "..."],
  "primary_endpoints": ["endpoint 1 with definition and assessment timepoint", "..."],
  "secondary_endpoints": ["endpoint 1 with definition and assessment timepoint", "..."],
  "estimated_enrollment": 100,
  "sample_size_justification": "Statistical rationale for sample size including assumptions, power, alpha, expected effect size",
  "statistical_analysis": "Primary analysis method, interim analyses, multiplicity adjustments, populations (ITT, per-protocol)",
  "safety_monitoring": "Description of Data Safety Monitoring Board (DSMB), stopping rules, dose-limiting toxicity definitions",
  "adverse_event_reporting": "AE/SAE grading criteria (e.g., CTCAE v5.0), reporting timelines, and expedited reporting requirements",
  "dose_modification": "Dose reduction levels, dose delay criteria, and discontinuation rules",
  "study_assessments": "Schedule of assessments: screening, on-treatment, end-of-treatment, and follow-up visit procedures (imaging, labs, PROs)",
  "study_schedule_table": [
    {"visit": "Screening", "timepoint": "Day -28 to -1", "procedures": "Informed consent, medical history, physical exam, labs, imaging, ECG, tumor biopsy"},
    {"visit": "Cycle 1 Day 1", "timepoint": "Day 1", "procedures": "Treatment administration, vitals, AE assessment"},
    {"visit": "Every 2 Cycles", "timepoint": "Every 8 weeks", "procedures": "Tumor assessment per RECIST 1.1, labs, AE assessment"},
    {"visit": "End of Treatment", "timepoint": "Within 30 days of last dose", "procedures": "Physical exam, labs, imaging, AE assessment"},
    {"visit": "Follow-up", "timepoint": "Every 3 months for 2 years", "procedures": "Survival status, subsequent therapies"}
  ],
  "ethical_considerations": "IRB/IEC approval requirements, Declaration of Helsinki compliance, informed consent process, vulnerable populations protections",
  "data_management": "EDC system, data quality assurance, source data verification, database lock procedures",
  "regulatory_considerations": "IND/CTA requirements, regulatory authority notifications, GCP compliance",
  "sex": "ALL or FEMALE",
  "minimum_age": "18 Years",
  "locations": [],
  "contact_info": "",
  "references": ["Key reference 1 (Author et al., Journal, Year)", "..."]
}

Base the protocol on patterns you see in the reference trials but make it a
coherent NEW study. Be specific, realistic, and scientifically rigorous.
Include specific drug names, dosing regimens, and measurable endpoints.
Follow ICH-GCP E6(R2) and FDA/EMA guidance for Phase II oncology trials.

IMPORTANT — administrative fields the user will fill in themselves:
- "protocol_id": return an empty string "". Do NOT invent a protocol ID.
- "sponsor": return an empty string "". Do NOT invent a sponsor.
- "locations": return an empty list []. Do NOT invent sites, cities, or countries.
- "contact_info": return an empty string "". Do NOT invent a PI or institution.
All other scientific/design fields should be fully drafted.

STUDY CENTER CONVENTION:
- Phase I and Phase II trials are typically SINGLE-CENTER (one academic/clinical
  site). Default to single-center for these unless the user explicitly asks for
  multi-center.
- Phase III trials and real-world evidence (RWE) studies are typically
  MULTI-CENTER.
Reflect this in "study_design" and "study_schema" (e.g., "Single-center,
randomized, open-label Phase II trial..."). The "locations" array still stays
empty — the user will supply the site.

OBSERVATIONAL / RWE ADAPTATIONS (when phase is "N/A"):
- NEVER use the words "Phase II" (or any phase) anywhere in title,
  official_title, description, study_design, study_schema, statistical_analysis,
  or any other field. This is not a phased trial.
- "study_design" should describe the observational design
  (e.g., "Multi-center retrospective cohort study", "Prospective observational
  registry", "Target trial emulation using claims data").
- "intervention_name" → describe the EXPOSURE being studied (e.g.,
  "CDK4/6 inhibitor therapy (palbociclib, ribociclib, or abemaciclib) plus
  endocrine therapy"), not a drug the investigator administers.
- "intervention_description" → describe how exposure is ascertained and
  categorized (EHR review, pharmacy claims, dispensing records), NOT dosing
  regimens the protocol prescribes.
- "comparator" → describe the reference cohort (e.g., "Endocrine therapy
  alone").
- "treatment_duration" → describe the observation / follow-up window.
- "primary_endpoints" / "secondary_endpoints" → real-world outcomes
  (rwPFS, rwOS, time-to-next-treatment, adherence, healthcare utilization).
- "safety_monitoring", "adverse_event_reporting", "dose_modification" →
  set to "Not applicable — observational study" (or describe passive
  pharmacovigilance / post-marketing surveillance signals).
- "study_assessments" and "study_schedule_table" → describe data collection
  timepoints and variables (index date, baseline covariates, follow-up
  outcomes), not on-treatment clinic visits.
- "sample_size_justification" → feasibility / precision for the effect size
  of interest, not power for an interventional endpoint.
- "statistical_analysis" → observational methods (propensity score matching /
  weighting, IPTW, Cox models with time-varying confounders, sensitivity
  analyses for unmeasured confounding, target trial emulation framework)."""


async def generate_protocol_json(
    llm: BaseLLMProvider,
    query: str,
    retrieved_docs: list[dict],
) -> dict:
    """Ask the LLM to draft a protocol based on RAG-retrieved trials; return parsed JSON dict."""
    context = format_context(retrieved_docs)

    user_msg = (
        f"The user wants to plan a breast cancer study with this focus:\n"
        f'"{query}"\n\n'
        f"=== REFERENCE TRIALS FROM DATABASE ===\n{context}\n"
        f"=== END ===\n\n"
        f"Infer the study type from the user's focus above (interventional "
        f"phase I-IV, or observational / real-world evidence). Draft the "
        f"protocol as JSON following the system prompt schema. If the user "
        f"asked for a real-world evidence / observational study, do NOT "
        f"label it as Phase II anywhere."
    )

    # Use streaming for long generations (max_tokens=8192). Non-streaming
    # requests with high max_tokens can exceed the API gateway timeout and
    # fail with APIConnectionError. Streaming keeps the connection warm
    # token-by-token; complete() accumulates the full response.
    raw = (
        await llm.complete(
            system=PROTOCOL_SYSTEM_PROMPT,
            messages=[{"role": "user", "content": user_msg}],
            max_tokens=8192,
        )
    ).strip()
    if raw.startswith("```"):
        raw = raw.split("\n", 1)[1]
        raw = raw.rsplit("```", 1)[0]

    protocol = json.loads(raw)
    logger.info("Protocol JSON generated with %d keys", len(protocol))
    return protocol


REFINE_SYSTEM_PROMPT = """\
You are revising an existing breast-cancer study protocol JSON. The request
may take one of three shapes:

(a) A single targeted edit ("change the comparator to pembrolizumab").
(b) A longer block of feedback pasted in from a peer reviewer or another LLM
    that lists multiple discrete changes ("Reviewer 1 says: 1) tighten
    inclusion criterion #3; 2) add an exclusion for prior anthracycline…").
(c) A short confirmation of changes you (the assistant) previously proposed
    in this conversation ("apply", "go ahead", "yes apply 1 and 3",
    "make these changes"). In this case, look back at your most recent
    assistant turn(s) in the refinement_messages for the proposed changes
    you advised, and apply exactly the subset the user confirmed:
       - "apply" / "apply all" / "go ahead" / "approved" / "do it" → apply
         every concrete change you had proposed.
       - "apply 1 and 3" / "apply items 2, 4" → apply only those numbered
         items.
       - "skip 2" / "all except 2" → apply everything proposed except item 2.
    If the conversation has no clearly-proposed set of changes (e.g. the
    user said "apply" with nothing to apply), return the protocol unchanged
    and explain in the NOTE: line.

When the request is multi-point feedback that hasn't been pre-discussed:
- Parse it into discrete actionable items first (don't conflate two reviewer
  comments into one change).
- Apply EACH item that is a clear, concrete edit. If an item is vague ("the
  endpoints don't feel right") or asks a question rather than requesting a
  change, leave that section alone and note it in the NOTE: line below.
- When two reviewers conflict (e.g. one says "broaden eligibility", another
  "tighten ECOG"), prefer the more conservative / safer option for clinical
  trial enrollment and call out the conflict in the NOTE: so the user can
  override.

Your job:
1. Apply the requested change(s) accurately and propagate any DIRECTLY IMPLIED
   downstream consistency fixes. For example, if the clinician corrects the
   control arm from "investigator's choice chemotherapy" to "pembrolizumab
   monotherapy", you must also fix:
     - primary_objectives / secondary_objectives that reference the comparator
     - intervention_description / study_schema / study_design arm descriptions
     - sample_size_justification if it assumed the wrong control effect size
     - any phrasing that miscategorizes the comparator (e.g. calling
       pembrolizumab a "chemotherapy")
2. Preserve EVERYTHING else exactly as-is. Do not rewrite unrelated sections,
   do not "improve" wording, do not add new fields, do not change administrative
   placeholders (protocol_id, sponsor, locations, contact_info stay as the
   user provided them — empty if empty).
3. Keep the same JSON schema (same keys, same types). Return the FULL updated
   protocol JSON object, not a diff.

After the JSON, on a new line, write a single line starting with "CHANGES:"
listing the keys you modified, comma-separated, e.g.
   CHANGES: comparator, primary_objectives, intervention_description

Then on the next line, write a single line starting with "NOTE:" with a short
plain-English summary of what changed and why, addressed to the clinician.
For multi-point feedback, structure the NOTE as a brief enumerated list so the
clinician can audit each change against the original feedback. Example:

   NOTE: Applied 4 of 5 feedback items. (1) Tightened inclusion criterion #3
   to require ECOG 0-1. (2) Added exclusion for prior anthracycline exposure
   within 12 months. (3) Updated sample size to 180 with revised power
   calculation. (4) Switched primary endpoint timepoint to 6 months.
   Reviewer 2's item about "considering an active comparator" was left for
   your decision — the current single-arm design is intentional for the
   Phase II setting; let me know if you want me to redesign as a 2-arm trial.

Output format (no markdown fences):
<full JSON object>
CHANGES: <comma-separated keys>
NOTE: <short summary or enumerated list of applied items>

Follow all the same domain conventions as the original protocol prompt:
- Empty admin fields stay empty.
- Observational/RWE studies keep phase "N/A" with no phase language anywhere.
- Single-center default for Phase I/II, multi-center for Phase III and RWE."""


async def refine_protocol_json(
    llm: BaseLLMProvider,
    current_protocol: dict,
    refinement_messages: list[dict],
    original_summary: str,
) -> tuple[dict, str, list[str]]:
    """Apply a clinician's correction request to an existing protocol.

    Returns (updated_protocol, assistant_note, changed_fields).
    `refinement_messages` is the running refinement conversation
    [{role, content}, ...] with the latest user request last.
    """
    history_lines = []
    for msg in refinement_messages[:-1]:
        role = "Clinician" if msg["role"] == "user" else "Assistant"
        history_lines.append(f"{role}: {msg['content']}")
    prior_history = "\n".join(history_lines) if history_lines else "(none)"

    latest_request = refinement_messages[-1]["content"] if refinement_messages else ""

    user_msg = (
        f"=== ORIGINAL PLANNING BRIEF ===\n{original_summary or '(not provided)'}\n\n"
        f"=== CURRENT PROTOCOL JSON ===\n{json.dumps(current_protocol, indent=2)}\n\n"
        f"=== PRIOR REFINEMENT TURNS ===\n{prior_history}\n\n"
        f"=== CLINICIAN'S CURRENT CORRECTION REQUEST ===\n{latest_request}\n\n"
        f"Apply the correction. Return the full updated JSON, then the "
        f"CHANGES: line, then the NOTE: line, per the system prompt."
    )

    raw = (
        await llm.complete(
            system=REFINE_SYSTEM_PROMPT,
            messages=[{"role": "user", "content": user_msg}],
            max_tokens=8192,
        )
    ).strip()
    if raw.startswith("```"):
        raw = raw.split("\n", 1)[1]
        raw = raw.rsplit("```", 1)[0].strip()

    # The model emits: <JSON object> \n CHANGES: ... \n NOTE: ...
    # Split off CHANGES: / NOTE: trailers, parse the JSON, keep the rest.
    note = ""
    changed: list[str] = []
    json_text = raw

    note_idx = raw.rfind("\nNOTE:")
    if note_idx != -1:
        note = raw[note_idx + len("\nNOTE:"):].strip()
        json_text = raw[:note_idx].rstrip()

    changes_idx = json_text.rfind("\nCHANGES:")
    if changes_idx != -1:
        changes_line = json_text[changes_idx + len("\nCHANGES:"):].strip()
        changed = [k.strip() for k in changes_line.split(",") if k.strip()]
        json_text = json_text[:changes_idx].rstrip()

    updated = json.loads(json_text)
    logger.info(
        "Protocol refined: %d changed fields (%s)",
        len(changed), ", ".join(changed) if changed else "none reported",
    )
    return updated, note, changed


def build_protocol_docx(protocol: dict) -> BytesIO:
    """Build a Word document from a protocol JSON dict. Returns a BytesIO buffer."""
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title page
    title = doc.add_heading(protocol.get("title", "Clinical Trial Protocol"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if protocol.get("official_title"):
        sub = doc.add_paragraph(protocol["official_title"])
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")
    # Administrative placeholders always render (Protocol ID / Sponsor are
    # intentionally left empty for the clinician to fill in).
    placeholder = "_" * 40
    meta_fields = [
        ("Protocol ID", "protocol_id", True),
        ("Sponsor", "sponsor", True),
        ("Phase", "phase", False),
        ("Status", "status", False),
        ("Conditions", "conditions", False),
    ]
    for label, key, always in meta_fields:
        value = protocol.get(key)
        if value:
            doc.add_paragraph(f"{label}: {value}")
        elif always:
            doc.add_paragraph(f"{label}: {placeholder}")

    doc.add_page_break()

    # Summary & Description
    _add_section(doc, "Summary", protocol.get("summary"))
    _add_section(doc, "Description", protocol.get("description"))

    # Objectives
    _add_list_section(doc, "Primary Objectives", protocol.get("primary_objectives"))
    _add_list_section(doc, "Secondary Objectives", protocol.get("secondary_objectives"))
    _add_list_section(doc, "Exploratory Objectives", protocol.get("exploratory_objectives"))

    # Study Design
    _add_section(doc, "Study Design", protocol.get("study_design"))
    _add_section(doc, "Study Schema", protocol.get("study_schema"))

    # Intervention
    _add_section(doc, "Intervention", protocol.get("intervention_name"))
    _add_section(doc, "Intervention Description", protocol.get("intervention_description"))
    _add_section(doc, "Comparator", protocol.get("comparator"))
    _add_section(doc, "Treatment Duration", protocol.get("treatment_duration"))

    # Eligibility
    _add_list_section(doc, "Inclusion Criteria", protocol.get("inclusion_criteria"))
    _add_list_section(doc, "Exclusion Criteria", protocol.get("exclusion_criteria"))

    # Endpoints
    _add_list_section(doc, "Primary Endpoints", protocol.get("primary_endpoints"))
    _add_list_section(doc, "Secondary Endpoints", protocol.get("secondary_endpoints"))

    # Statistics
    enrollment = protocol.get("estimated_enrollment")
    if enrollment:
        _add_section(doc, "Estimated Enrollment", str(enrollment))
    _add_section(doc, "Sample Size Justification", protocol.get("sample_size_justification"))
    _add_section(doc, "Statistical Analysis", protocol.get("statistical_analysis"))

    # Safety
    _add_section(doc, "Safety Monitoring", protocol.get("safety_monitoring"))
    _add_section(doc, "Adverse Event Reporting", protocol.get("adverse_event_reporting"))
    _add_section(doc, "Dose Modification", protocol.get("dose_modification"))

    # Assessments
    _add_section(doc, "Study Assessments", protocol.get("study_assessments"))

    # Schedule table
    schedule = protocol.get("study_schedule_table")
    if schedule and isinstance(schedule, list):
        doc.add_heading("Study Schedule", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        for i, header in enumerate(["Visit", "Timepoint", "Procedures"]):
            table.rows[0].cells[i].text = header
        for row_data in schedule:
            row = table.add_row()
            row.cells[0].text = row_data.get("visit", "")
            row.cells[1].text = row_data.get("timepoint", "")
            row.cells[2].text = row_data.get("procedures", "")

    # Regulatory & Ethics
    _add_section(doc, "Ethical Considerations", protocol.get("ethical_considerations"))
    _add_section(doc, "Data Management", protocol.get("data_management"))
    _add_section(doc, "Regulatory Considerations", protocol.get("regulatory_considerations"))

    # Demographics
    if protocol.get("sex"):
        _add_section(doc, "Eligible Sex", protocol["sex"])
    if protocol.get("minimum_age"):
        _add_section(doc, "Minimum Age", protocol["minimum_age"])

    # Locations — always render with a placeholder when the clinician
    # has not yet filled in study sites.
    locations = protocol.get("locations")
    if locations:
        _add_list_section(doc, "Locations", locations)
    else:
        doc.add_heading("Locations", level=1)
        doc.add_paragraph(placeholder)
        doc.add_paragraph("(to be completed: site name, city, state, country)")

    # Contact — always render with a placeholder.
    contact = protocol.get("contact_info")
    doc.add_heading("Contact Information", level=1)
    if contact:
        doc.add_paragraph(contact)
    else:
        doc.add_paragraph(f"Principal Investigator: {placeholder}")
        doc.add_paragraph(f"Institution: {placeholder}")
        doc.add_paragraph(f"Email: {placeholder}")
        doc.add_paragraph(f"Phone: {placeholder}")

    # References
    _add_list_section(doc, "References", protocol.get("references"))

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _add_section(doc: Document, title: str, content: str | None):
    if not content:
        return
    doc.add_heading(title, level=1)
    doc.add_paragraph(content)


def _add_list_section(doc: Document, title: str, items: list | None):
    if not items:
        return
    doc.add_heading(title, level=1)
    for item in items:
        doc.add_paragraph(str(item), style="List Bullet")


# ---------------------------------------------------------------------------
# PDF builder (reportlab / platypus)
# ---------------------------------------------------------------------------

def _pdf_styles():
    """Paragraph styles for the protocol PDF."""
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "ProtocolTitle", parent=base["Title"], fontSize=20, leading=24,
            alignment=TA_CENTER, spaceAfter=6,
        ),
        "subtitle": ParagraphStyle(
            "ProtocolSubtitle", parent=base["Italic"], fontSize=12, leading=15,
            alignment=TA_CENTER, spaceAfter=18,
        ),
        "meta": ParagraphStyle(
            "ProtocolMeta", parent=base["Normal"], fontSize=11, leading=15,
            spaceAfter=4,
        ),
        "h1": ParagraphStyle(
            "ProtocolH1", parent=base["Heading1"], fontSize=14, leading=18,
            spaceBefore=14, spaceAfter=6, textColor=colors.HexColor("#1a73e8"),
        ),
        "body": ParagraphStyle(
            "ProtocolBody", parent=base["BodyText"], fontSize=10.5, leading=14,
            spaceAfter=6,
        ),
        "bullet": ParagraphStyle(
            "ProtocolBullet", parent=base["BodyText"], fontSize=10.5, leading=14,
            leftIndent=12,
        ),
        "placeholder": ParagraphStyle(
            "ProtocolPlaceholder", parent=base["Italic"], fontSize=10.5, leading=14,
            textColor=colors.HexColor("#666666"), spaceAfter=4,
        ),
    }


def _pdf_escape(text) -> str:
    """Escape text for reportlab's Paragraph mini-markup (&, <, >)."""
    if text is None:
        return ""
    s = str(text)
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _pdf_section(flow, styles, title: str, content: str | None):
    if not content:
        return
    flow.append(Paragraph(title, styles["h1"]))
    flow.append(Paragraph(_pdf_escape(content), styles["body"]))


def _pdf_list_section(flow, styles, title: str, items: list | None):
    if not items:
        return
    flow.append(Paragraph(title, styles["h1"]))
    flow.append(ListFlowable(
        [ListItem(Paragraph(_pdf_escape(i), styles["bullet"])) for i in items],
        bulletType="bullet", leftIndent=18, bulletFontSize=10,
    ))
    flow.append(Spacer(1, 6))


def build_protocol_pdf(protocol: dict) -> BytesIO:
    """Build a PDF document from a protocol JSON dict. Returns a BytesIO buffer.

    Mirrors build_protocol_docx section-for-section, including the visible
    placeholder behaviour for Protocol ID, Sponsor, Locations, and
    Contact Information when the clinician hasn't filled them in yet.
    """
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=LETTER,
        leftMargin=0.9 * inch, rightMargin=0.9 * inch,
        topMargin=0.8 * inch, bottomMargin=0.8 * inch,
        title=protocol.get("title", "Clinical Trial Protocol"),
    )
    styles = _pdf_styles()
    flow = []

    # Title page
    flow.append(Paragraph(_pdf_escape(protocol.get("title", "Clinical Trial Protocol")), styles["title"]))
    if protocol.get("official_title"):
        flow.append(Paragraph(_pdf_escape(protocol["official_title"]), styles["subtitle"]))

    placeholder = "_" * 40
    meta_fields = [
        ("Protocol ID", "protocol_id", True),
        ("Sponsor", "sponsor", True),
        ("Phase", "phase", False),
        ("Status", "status", False),
        ("Conditions", "conditions", False),
    ]
    for label, key, always in meta_fields:
        value = protocol.get(key)
        if value:
            flow.append(Paragraph(f"<b>{_pdf_escape(label)}:</b> {_pdf_escape(value)}", styles["meta"]))
        elif always:
            flow.append(Paragraph(f"<b>{_pdf_escape(label)}:</b> {placeholder}", styles["meta"]))

    flow.append(PageBreak())

    # Summary & Description
    _pdf_section(flow, styles, "Summary", protocol.get("summary"))
    _pdf_section(flow, styles, "Description", protocol.get("description"))

    # Objectives
    _pdf_list_section(flow, styles, "Primary Objectives", protocol.get("primary_objectives"))
    _pdf_list_section(flow, styles, "Secondary Objectives", protocol.get("secondary_objectives"))
    _pdf_list_section(flow, styles, "Exploratory Objectives", protocol.get("exploratory_objectives"))

    # Study Design
    _pdf_section(flow, styles, "Study Design", protocol.get("study_design"))
    _pdf_section(flow, styles, "Study Schema", protocol.get("study_schema"))

    # Intervention
    _pdf_section(flow, styles, "Intervention", protocol.get("intervention_name"))
    _pdf_section(flow, styles, "Intervention Description", protocol.get("intervention_description"))
    _pdf_section(flow, styles, "Comparator", protocol.get("comparator"))
    _pdf_section(flow, styles, "Treatment Duration", protocol.get("treatment_duration"))

    # Eligibility
    _pdf_list_section(flow, styles, "Inclusion Criteria", protocol.get("inclusion_criteria"))
    _pdf_list_section(flow, styles, "Exclusion Criteria", protocol.get("exclusion_criteria"))

    # Endpoints
    _pdf_list_section(flow, styles, "Primary Endpoints", protocol.get("primary_endpoints"))
    _pdf_list_section(flow, styles, "Secondary Endpoints", protocol.get("secondary_endpoints"))

    # Statistics
    enrollment = protocol.get("estimated_enrollment")
    if enrollment:
        _pdf_section(flow, styles, "Estimated Enrollment", str(enrollment))
    _pdf_section(flow, styles, "Sample Size Justification", protocol.get("sample_size_justification"))
    _pdf_section(flow, styles, "Statistical Analysis", protocol.get("statistical_analysis"))

    # Safety
    _pdf_section(flow, styles, "Safety Monitoring", protocol.get("safety_monitoring"))
    _pdf_section(flow, styles, "Adverse Event Reporting", protocol.get("adverse_event_reporting"))
    _pdf_section(flow, styles, "Dose Modification", protocol.get("dose_modification"))

    # Assessments
    _pdf_section(flow, styles, "Study Assessments", protocol.get("study_assessments"))

    # Schedule table
    schedule = protocol.get("study_schedule_table")
    if schedule and isinstance(schedule, list):
        flow.append(Paragraph("Study Schedule", styles["h1"]))
        data = [["Visit", "Timepoint", "Procedures"]]
        for row in schedule:
            data.append([
                Paragraph(_pdf_escape(row.get("visit", "")), styles["body"]),
                Paragraph(_pdf_escape(row.get("timepoint", "")), styles["body"]),
                Paragraph(_pdf_escape(row.get("procedures", "")), styles["body"]),
            ])
        # Total content width = page - margins = 8.5" - 1.8" = 6.7"
        tbl = Table(data, colWidths=[1.4 * inch, 1.4 * inch, 3.9 * inch])
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8f0fe")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1a73e8")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
            ("TOPPADDING", (0, 0), (-1, 0), 6),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        flow.append(tbl)
        flow.append(Spacer(1, 8))

    # Regulatory & Ethics
    _pdf_section(flow, styles, "Ethical Considerations", protocol.get("ethical_considerations"))
    _pdf_section(flow, styles, "Data Management", protocol.get("data_management"))
    _pdf_section(flow, styles, "Regulatory Considerations", protocol.get("regulatory_considerations"))

    # Demographics
    if protocol.get("sex"):
        _pdf_section(flow, styles, "Eligible Sex", protocol["sex"])
    if protocol.get("minimum_age"):
        _pdf_section(flow, styles, "Minimum Age", protocol["minimum_age"])

    # Locations — always render with a placeholder when empty
    locations = protocol.get("locations")
    if locations:
        _pdf_list_section(flow, styles, "Locations", locations)
    else:
        flow.append(Paragraph("Locations", styles["h1"]))
        flow.append(Paragraph(placeholder, styles["placeholder"]))
        flow.append(Paragraph(
            "(to be completed: site name, city, state, country)", styles["placeholder"],
        ))

    # Contact — always render with a placeholder when empty
    contact = protocol.get("contact_info")
    flow.append(Paragraph("Contact Information", styles["h1"]))
    if contact:
        flow.append(Paragraph(_pdf_escape(contact), styles["body"]))
    else:
        for sub_label in ("Principal Investigator", "Institution", "Email", "Phone"):
            flow.append(Paragraph(
                f"<b>{sub_label}:</b> {placeholder}", styles["placeholder"],
            ))

    # References
    _pdf_list_section(flow, styles, "References", protocol.get("references"))

    doc.build(flow)
    buffer.seek(0)
    return buffer
